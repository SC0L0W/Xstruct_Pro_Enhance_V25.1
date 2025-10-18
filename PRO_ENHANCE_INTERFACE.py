import tkinter as tk
import threading
import time
import os
import sys
import openpyxl
import logging
import webbrowser
import math
from tkinter import Canvas, Entry, Button, PhotoImage, Frame, filedialog, OptionMenu, StringVar, IntVar, messagebox, Toplevel, Label
from PIL import Image, ImageTk, ImageSequence
from openpyxl import Workbook, load_workbook
from openstaad import Root, Geometry, Properties, Output



class Interface(tk.Toplevel):
    def __init__(self, master=None):
        super().__init__(master)
        self.setup_window()
        self.load_images()
        self.create_menu_bar()
        self.create_pages()


    def setup_window(self):
        """Set up the main window."""
        self.geometry("1080x720")
        self.title("PRO_ENHANCE")
        self.resizable(False, False)

        # Center the window on the screen
        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()
        position_top = int(screen_height / 2 - 720 / 2)
        position_left = int(screen_width / 2 - 1080 / 2)
        self.geometry(f"1080x720+{position_left}+{position_top}")

        # Set the window icon
        icon_image = Image.open(self.get_resource_path("Icon.ico"))
        self.iconphoto(False, ImageTk.PhotoImage(icon_image))

    def load_images(self):
        """Load required images for buttons."""
        self.toggle_icon = tk.PhotoImage(file=self.get_resource_path("images_sidebar/toggle_btn_icon.png"))
        self.staad_icon = tk.PhotoImage(file=self.get_resource_path("images_sidebar/staad_logo.png"))
        self.concrete_icon = tk.PhotoImage(file=self.get_resource_path("images_sidebar/concrete.png"))
        self.steel_icon = tk.PhotoImage(file=self.get_resource_path("images_sidebar/steel.png"))
        self.home_icon = tk.PhotoImage(file=self.get_resource_path("images_sidebar/home_icon.png"))
        self.close_icon = tk.PhotoImage(file=self.get_resource_path("images_sidebar/close_btn_icon.png"))

    def create_menu_bar(self):
        """Create the menu bar on the left."""
        self.menu_bar_colour = "#0067A0"
        self.menu_bar_frame = tk.Frame(self, bg=self.menu_bar_colour)
        self.menu_bar_frame.pack(side=tk.LEFT, fill=tk.Y, pady=4, padx=4)
        self.menu_bar_frame.configure(width=60)

        # Toggle Menu Button
        self.toggle_menu_btn = tk.Button(
            self.menu_bar_frame,
            image=self.toggle_icon,
            bg=self.menu_bar_colour,
            bd=0,
            activebackground=self.menu_bar_colour,
            command=self.extend_menu_bar,
        )
        self.toggle_menu_btn.place(x=10, y=10)

        # STAAD Button
        self.staad_btn = tk.Button(
            self.menu_bar_frame,
            image=self.staad_icon,
            bg=self.menu_bar_colour,
            bd=0,
            activebackground=self.menu_bar_colour,
            command=lambda: self.switch_indication(self.staad_btn_indicator, self.staad_page),
        )
        self.staad_btn.place(x=10, y=130, width=40, height=40)

        self.staad_btn_indicator = tk.Label(self.menu_bar_frame, bg=self.menu_bar_colour)
        self.staad_btn_indicator.place(x=5, y=130, width=5, height=40)

        # STAAD Label
        self.staad_page_lb = tk.Label(
            self.menu_bar_frame,
            text="STAAD",
            bg=self.menu_bar_colour,
            fg="White",
            font=("Bold", 12),
            anchor=tk.W,
        )
        self.staad_page_lb.place(x=60, y=130, width=150, height=40)
        self.staad_page_lb.bind('<Button-1>',
                                lambda e: self.switch_indication(self.staad_btn_indicator, self.staad_page))


        # Concrete Button
        self.concrete_btn = tk.Button(
            self.menu_bar_frame,
            image=self.concrete_icon,
            bg=self.menu_bar_colour,
            bd=0,
            activebackground=self.menu_bar_colour,
            command=lambda: self.switch_indication(self.concrete_btn_indicator, self.concrete_page),
        )
        self.concrete_btn.place(x=10, y=190, width=40, height=40)

        self.concrete_btn_indicator = tk.Label(self.menu_bar_frame, bg=self.menu_bar_colour)
        self.concrete_btn_indicator.place(x=5, y=190, width=5, height=40)

        self.concrete_btn_lb = tk.Label(
            self.menu_bar_frame,
            text="CONCRETE",
            bg=self.menu_bar_colour,
            fg="White",
            font=("Bold", 12),
            anchor=tk.W,
        )
        self.concrete_btn_lb.place(x=60, y=190, width=150, height=40)
        self.concrete_btn_lb.bind('<Button-1>',
                                lambda e: self.switch_indication(self.concrete_btn_indicator, self.concrete_page))

        # Steel Button
        self.steel_btn = tk.Button(
            self.menu_bar_frame,
            image=self.steel_icon,
            bg=self.menu_bar_colour,
            bd=0,
            activebackground=self.menu_bar_colour,
            command=lambda: self.switch_indication(self.steel_btn_indicator, self.steel_page),
        )
        self.steel_btn.place(x=10, y=250, width=40, height=40)

        self.steel_btn_indicator = tk.Label(self.menu_bar_frame, bg=self.menu_bar_colour)
        self.steel_btn_indicator.place(x=5, y=250, width=5, height=40)

        self.steel_btn_lb = tk.Label(
            self.menu_bar_frame,
            text="STEEL",
            bg=self.menu_bar_colour,
            fg="White",
            font=("Bold", 12),
            anchor=tk.W,
        )
        self.steel_btn_lb.place(x=60, y=250, width=150, height=40)

        self.steel_btn_lb.bind('<Button-1>',
                                  lambda e: self.switch_indication(self.steel_btn_indicator, self.steel_page))

        self.home_btn = tk.Button(
            self.menu_bar_frame,
            image=self.home_icon,
            bg=self.menu_bar_colour,
            bd=0,
            activebackground=self.menu_bar_colour,
            command=self.show_loading_screen  # Pass the method reference, not the result of calling it
        )

        self.home_lb = tk.Label(
            self.menu_bar_frame,
            text="HOME",
            bg=self.menu_bar_colour,
            fg="White",
            font=("Bold", 12),
            anchor=tk.W,
        )
        self.home_btn.place(x=10, y=670, width=40, height=40)
        self.home_lb.place(x=60, y=670, width=150, height=40)

        # Bind the label click event correctly by passing the method reference
        self.home_lb.bind('<Button-1>', lambda event: self.show_loading_screen())

    def switch_indication(self, indicator_lb, page):
        """Switch page and indicate the selected page in the menu."""
        # Reset all indicators to the default background color
        self.staad_btn_indicator.config(bg=self.menu_bar_colour)
        self.concrete_btn_indicator.config(bg=self.menu_bar_colour)
        self.steel_btn_indicator.config(bg=self.menu_bar_colour)

        # Highlight the selected page's indicator
        indicator_lb.config(bg='white')

        # Fold the menu bar if it's extended
        if self.menu_bar_frame.winfo_width() > 45:
            self.fold_menu_bar()

        # Clear the current page content
        for widget in self.page_frame.winfo_children():
            widget.destroy()

        # Render the new page
        page()

    def switch_page(self, page_function):
        """Switch to the selected page."""
        for widget in self.page_frame.winfo_children():
            widget.destroy()  # Clear the current page content
        page_function()  # Render the selected page

    def create_pages(self):
        """Create the placeholders for pages."""
        self.page_frame = tk.Frame(self)
        self.page_frame.place(relwidth=1.0, relheight=1.0, x=70)

        # Example: Initialize STAAD page
        self.staad_page()

    def staad_page(self):
        logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

        staad_page_fm = Frame(self.page_frame)
        staad_page_fm.place(relwidth=1.0, relheight=1.0)

        canvas = Canvas(staad_page_fm, bg="#FFFFFF", height=720, width=1016, bd=0, highlightthickness=0)
        canvas.place(x=0, y=0)
        canvas.create_rectangle(0, 0, 1016, 720, fill="#FFFFFF", outline="")

        canvas1 = Canvas(staad_page_fm, bg="#FFFFFF", height=720, width=1016, bd=0, highlightthickness=0)
        canvas1.create_rectangle(0, 0, 1016, 720, fill="#FFFFFF", outline="")

        image_image_1 = PhotoImage(file=self.get_resource_path("images_staadpage/image_1.png"))
        canvas.image_1 = image_image_1
        canvas.create_image(501.0, 396.0, image=image_image_1)
        canvas1.create_image(501.0, 396.0, image=image_image_1)

        image_image_7 = PhotoImage(file=self.get_resource_path("images_staadpage/image_7.png"))
        canvas.image_image_7 = image_image_7
        canvas.create_image(500.0, 708.0, image=image_image_7)
        canvas1.create_image(500.0, 708.0, image=image_image_7)

        image_image_9 = PhotoImage(file=self.get_resource_path("images_staadpage/image_9.png"))
        canvas.image_image_9 = image_image_9
        canvas.create_image(500, 24.0, image=image_image_9)
        canvas1.create_image(500, 24.0, image=image_image_9)

        image_image_2 = PhotoImage(file=self.get_resource_path("images_staadpage/image_2.png"))
        canvas.image_2 = image_image_2
        canvas.create_image(230.0, 424.0, image=image_image_2)

        image_image_3 = PhotoImage(file=self.get_resource_path("images_staadpage/image_3.png"))
        canvas.image_3 = image_image_3
        canvas.create_image(230.0, 554.0, image=image_image_3)

        image_image_4 = PhotoImage(file=self.get_resource_path("images_staadpage/image_4.png"))
        canvas.image_4 = image_image_4
        canvas.create_image(230.0, 121.0, image=image_image_4)

        image_image_5 = PhotoImage(file=self.get_resource_path("images_staadpage/image_5.png"))
        canvas.image_image_5 = image_image_5
        canvas.create_image(238.0, 298.0, image=image_image_5)

        image_image_6 = PhotoImage(file=self.get_resource_path("images_staadpage/image_6.png"))
        canvas.image_image_6 = image_image_6
        canvas.create_image(70.0, 171.0, image=image_image_6)

        image_image_8 = PhotoImage(file=self.get_resource_path("images_staadpage/image_8.png"))
        canvas.image_8 = image_image_8
        canvas.create_image(70.0, 585.0, image=image_image_8)

        entry_filename = Entry(canvas,bd=0, bg="#FFFFFF", fg="#000716", highlightthickness=0)
        canvas.create_window(114.0, 146.0, anchor="nw", window=entry_filename, width=300.0, height=18.0)
        entry_file_location = Entry(canvas,bd=0, bg="#FFFFFF", fg="#000716", highlightthickness=0)
        canvas.create_window(114.0, 176.0, anchor="nw", window=entry_file_location, width=300.0, height=18.0)

        def choose_location():
            folder_path = filedialog.askdirectory()
            entry_file_location.delete(0, 'end')
            entry_file_location.insert(0, folder_path)

        entry_file_location.bind("<Button-1>", lambda e: choose_location())

        def create_excel_file():
            filename = entry_filename.get()
            folder_location = entry_file_location.get()

            if not filename or not folder_location:
                messagebox.showwarning("Input Error", "Please provide both filename and folder location.")
                return

            full_path = os.path.join(folder_location, filename + ".xlsx")
            wb = Workbook()
            wb.save(full_path)
            messagebox.showinfo("Success", f"PRO ENHANCE file created at: {full_path}")

            entry_xstru_excelfile.delete(0, 'end')
            entry_xstru_excelfile.insert(0, filename + ".xlsx")

        button_create = PhotoImage(file=self.get_resource_path("images_staadpage/button_4.png"))
        canvas.button_image_4 = button_create
        button_create = Button(canvas,
            image=button_create,
            borderwidth=0,
            highlightthickness=0,
            command=create_excel_file,
            bg="green",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(115, 210, anchor="nw", window=button_create, width=72.0, height=20.0)

        class Beam3D:
            def __init__(self, canvas, excel_path):
                self.canvas = canvas
                self.width = canvas.winfo_width()
                self.height = canvas.winfo_height()
                self.excel_path = excel_path
                self.beams = []
                self.scale = 0.7
                self.offset_x = self.width / 2
                self.offset_y = self.height / 2
                self.is_dragging = False
                self.is_panning = False
                self.angle_x = math.radians(5)
                self.angle_y = math.radians(5)

                self.bounding_box = [470.0, 100.0, 1000.0, 690.0]
                self.bounding_width = self.bounding_box[2] - self.bounding_box[0]
                self.bounding_height = self.bounding_box[3] - self.bounding_box[1]

                self.last_x = 0
                self.last_y = 0

                self.update_canvas_size()

            def update_canvas_size(self, event=None):
                self.width = self.canvas.winfo_width()
                self.height = self.canvas.winfo_height()
                self.offset_x = self.width / 2
                self.offset_y = self.height / 2

            def load_beam_data(self):
                if not self.excel_path:
                    messagebox.showerror("File Selection Error", "No file selected. Please select an Excel file.")
                    return
                try:
                    workbook = openpyxl.load_workbook(self.excel_path)
                    sheet = workbook.active
                    self.beams.clear()

                    for row in sheet.iter_rows(min_row=2):
                        length = row[1].value
                        section = row[2].value
                        start_x, start_y, start_z = row[6].value, row[7].value, row[8].value
                        end_x, end_y, end_z = row[9].value, row[10].value, row[11].value

                        if section.startswith("Rect"):
                            dims = section.split()[1].split("x")
                            width, height = float(dims[0]), float(dims[1])
                            shape = "rect"
                        elif section.startswith("Cir"):
                            diameter = float(section.split()[1])
                            width, height = diameter, diameter
                            shape = "circle"

                        self.beams.append({
                            "length": length,
                            "shape": shape,
                            "width": width,
                            "height": height,
                            "start": (start_x, start_y, start_z),
                            "end": (end_x, end_y, end_z)
                        })

                    self.adjust_initial_scale()


                except openpyxl.utils.exceptions.InvalidFileException:
                    messagebox.showerror("File Format Error", "The selected file format is not supported by openpyxl.")
                except FileNotFoundError:
                    messagebox.showerror("File Not Found", "The file could not be found. Please check the file path.")
                except Exception as e:
                    messagebox.showerror("Error", f"Error loading beam data: {e}")

            def adjust_initial_scale(self):
                if not self.beams:
                    return
                min_x = min(beam["start"][0] for beam in self.beams)
                max_x = max(beam["end"][0] for beam in self.beams)
                min_y = min(beam["start"][1] for beam in self.beams)
                max_y = max(beam["end"][1] for beam in self.beams)
                min_z = min(beam["start"][2] for beam in self.beams)
                max_z = max(beam["end"][2] for beam in self.beams)

                width_range = max_x - min_x
                height_range = max_y - min_y
                depth_range = max_z - min_z

                max_range = max(width_range, height_range, depth_range)

                self.scale = (min(self.bounding_width, self.bounding_height) / max_range)*0.7

                center_x = (min_x + max_x) / 0.57
                center_y = (min_y + max_y) / 2.5

                self.offset_x = self.bounding_box[0] + self.bounding_width / 2 - center_x * self.scale
                self.offset_y = self.bounding_box[1] + self.bounding_height / 2 - center_y * self.scale

            def project(self, x, y, z):
                cos_x = math.cos(self.angle_x)
                sin_x = math.sin(self.angle_x)
                cos_y = math.cos(self.angle_y)
                sin_y = math.sin(self.angle_y)

                y, z = y * cos_x - z * sin_x, y * sin_x + z * cos_x

                x, z = x * cos_y + z * sin_y, -x * sin_y + z * cos_y

                scale = self.scale
                projected_x = x * scale
                projected_y = -y * scale
                final_x = projected_x + self.offset_x
                final_y = projected_y + self.offset_y

                return (final_x, final_y)

            def draw_beams(self):
                self.canvas.delete("all")
                for beam in self.beams:
                    start_x, start_y, start_z = beam["start"]
                    end_x, end_y, end_z = beam["end"]

                    start_2d = self.project(start_x, start_y, start_z)
                    end_2d = self.project(end_x, end_y, end_z)

                    self.canvas.create_line(start_2d, end_2d, fill="Black", width=2)

                self.draw_axes()

                for beam in beam_model.beams:
                    start_2d = beam_model.project(*beam["start"])
                    end_2d = beam_model.project(*beam["end"])
                    beam_model.draw_section(beam, start_2d, end_2d)

            def draw_section(self, beam, start_2d, end_2d):
                width = beam["width"] * self.scale / 2  # Half width for scaling
                if beam["shape"] == "rect":
                    x1_start, y1_start = start_2d[0] - width, start_2d[1] - width
                    x2_start, y2_start = start_2d[0] + width, start_2d[1] + width
                    x1_end, y1_end = end_2d[0] - width, end_2d[1] - width
                    x2_end, y2_end = end_2d[0] + width, end_2d[1] + width

                    self.canvas.create_rectangle(x1_start, y1_start, x2_start, y2_start, outline="blue", fill="#ADD8E6")
                    self.canvas.create_rectangle(x1_end, y1_end, x2_end, y2_end, outline="blue", fill="#5F9EA0")

                    self.canvas.create_line(x1_start, y1_start, x1_end, y1_end, fill="#4682B4")
                    self.canvas.create_line(x2_start, y1_start, x2_end, y1_end, fill="#4682B4")
                    self.canvas.create_line(x1_start, y2_start, x1_end, y2_end, fill="#4682B4")
                    self.canvas.create_line(x2_start, y2_start, x2_end, y2_end, fill="#4682B4")

                elif beam["shape"] == "circle":
                    x1_start, y1_start = start_2d[0] - width, start_2d[1] - width
                    x2_start, y2_start = start_2d[0] + width, start_2d[1] + width
                    x1_end, y1_end = end_2d[0] - width, end_2d[1] - width
                    x2_end, y2_end = end_2d[0] + width, end_2d[1] + width

                    self.canvas.create_oval(x1_start, y1_start, x2_start, y2_start, outline="red", fill="#FF6347")
                    self.canvas.create_oval(x1_end, y1_end, x2_end, y2_end, outline="red", fill="#CD5C5C")

                    self.canvas.create_line((x1_start + x2_start) / 2, y1_start, (x1_end + x2_end) / 2, y1_end,
                                            fill="#A52A2A")
                    self.canvas.create_line((x1_start + x2_start) / 2, y2_start, (x1_end + x2_end) / 2, y2_end,
                                            fill="#A52A2A")

            def draw_axes(self):
                self.canvas.create_line(self.project(-3, 0, 0), self.project(10, 0, 0), fill="blue", width=2,
                                        arrow=tk.LAST)
                self.canvas.create_text(self.project(10, 0, 0), text="X", fill="blue")

                self.canvas.create_line(self.project(0, -3, 0), self.project(0, 10, 0), fill="red", width=2,
                                        arrow=tk.LAST)
                self.canvas.create_text(self.project(0, 10, 0), text="Y", fill="red")

                self.canvas.create_line(self.project(0, 0, -3), self.project(0, 0, 10), fill="green", width=2,
                                        arrow=tk.LAST)
                self.canvas.create_text(self.project(0, 0, 10), text="Z", fill="green")

            def start_horizontal_rotation(self, event):
                self.last_x = event.x
                self.last_y = event.y

            def start_vertical_rotation(self, event):
                self.last_x = event.x
                self.last_y = event.y

            def rotate_horizontal(self, event):
                dx = event.x - self.last_x
                self.angle_y -= dx * 0.005
                self.last_x = event.x
                self.draw_beams()

            def rotate_vertical(self, event):
                dy = event.y - self.last_y
                self.angle_x += dy * 0.005
                self.last_y = event.y
                self.draw_beams()

            def on_drag_start(self, event):
                self.is_dragging = True
                self.last_x = event.x
                self.last_y = event.y

            def on_drag_motion(self, event):
                if self.is_dragging:
                    dx = event.x - self.last_x
                    dy = event.y - self.last_y
                    self.offset_x += dx
                    self.offset_y += dy
                    self.last_x = event.x
                    self.last_y = event.y
                    self.draw_beams()

            def on_drag_end(self, event):
                self.is_dragging = False

            def zoom(self, event):
                if event.delta > 0:
                    self.scale *= 1.1
                else:
                    self.scale /= 1.1
                self.draw_beams()

            def set_isometric_view(self):
                self.angle_x = math.radians(10)
                self.angle_y = math.radians(10)
                self.angle_z = math.radians(0)
                self.draw_beams()

            def set_top_view(self):
                self.angle_x = math.radians(90)
                self.angle_y = 0
                self.angle_z = 0
                self.draw_beams()

            def set_left_view(self):
                self.angle_x = 0
                self.angle_y = math.radians(90)
                self.angle_z = 0
                self.draw_beams()

            def set_front_view(self):
                self.angle_x = 0
                self.angle_y = 0
                self.angle_z = 0
                self.draw_beams()

        canvas3d = tk.Canvas(canvas, width=530, height=590, bg="#FFFFFF", highlightthickness=0)
        canvas.create_window(470.0, 100.0, anchor="nw", window=canvas3d)
        beam_model = Beam3D(canvas3d, "")

        def on_load_beams():
            excel_path = entry_xstru_excelfile.get()
            beam_model.excel_path = excel_path
            beam_model.load_beam_data()
            beam_model.draw_beams()

        canvas3d.bind("<ButtonPress-1>", beam_model.start_horizontal_rotation)
        canvas3d.bind("<B1-Motion>", beam_model.rotate_horizontal)
        canvas3d.bind("<ButtonPress-3>", beam_model.start_vertical_rotation)
        canvas3d.bind("<B3-Motion>", beam_model.rotate_vertical)
        canvas3d.bind("<MouseWheel>", beam_model.zoom)
        canvas3d.bind("<ButtonPress-2>", beam_model.on_drag_start)
        canvas3d.bind("<B2-Motion>", beam_model.on_drag_motion)
        canvas3d.bind("<ButtonRelease-2>", beam_model.on_drag_end)

        isometric_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/IsometricView.png"))
        canvas.isometric_button = isometric_button
        isometric_button = tk.Button(
            canvas,
            image=isometric_button,
            borderwidth=0,
            highlightthickness=0,
            command=beam_model.set_isometric_view,
            bg="White",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(970, 625, anchor="nw", window=isometric_button, width=30.0, height=30.0)

        Top_View_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/TopView.png"))
        canvas.Top_View_button = Top_View_button
        Top_View_button = tk.Button(
            canvas,
            image=Top_View_button,
            borderwidth=0,
            highlightthickness=0,
            command=beam_model.set_top_view,
            bg="White",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(940, 595, anchor="nw", window=Top_View_button, width=30.0, height=30.0)

        Left_View_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/LeftView.png"))
        canvas.Left_View_button = Left_View_button
        Left_View_button = tk.Button(
            canvas,
            image=Left_View_button,
            borderwidth=0,
            highlightthickness=0,
            command=beam_model.set_left_view,
            bg="White",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(910, 625, anchor="nw", window=Left_View_button, width=30.0, height=30.0)

        Front_View_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/FrontView.png"))
        canvas.Front_View_button = Front_View_button
        Front_View_button = tk.Button(
            canvas,
            image=Front_View_button,
            borderwidth=0,
            highlightthickness=0,
            command=beam_model.set_front_view,
            bg="White",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(940, 655, anchor="nw", window=Front_View_button, width=30.0, height=30.0)

        Axis_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/AxisArrows.png"))
        canvas.Axis_button = Axis_button
        Axis_button = tk.Button(
            canvas,
            image=Axis_button,
            borderwidth=0,
            highlightthickness=0,
            command=beam_model.draw_beams(),
            bg="White",
            activebackground="White",
            relief="flat"
        )
        canvas.create_window(941, 625, anchor="nw", window=Axis_button, width=30.0, height=30.0)
        # Add instructions text

        instruction_text = [
            "Left-click and drag: Horizontal rotation",
            "Right-click and drag: Vertical rotation",
            "Scroll wheel: Zoom in/out",
            "Middle-click and drag: Pan"
        ]

        instruction_button_image = tk.PhotoImage(file=self.get_resource_path("images_staadpage/Help.png"))
        canvas.instruction_button = instruction_button_image
        instruction_button = tk.Button(
            canvas,
            image=instruction_button_image,
            borderwidth=0,
            highlightthickness=0,
            command=on_load_beams,
            bg="White",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(476, 656, anchor="nw", window=instruction_button, width=30.0, height=30.0)

        instruction_canvas = tk.Canvas(canvas, width=300, height=100, bg="lightyellow", highlightthickness=0)

        def show_instructions(event):
            instruction_canvas.place(x=510, y=570)
            y_offset = 10
            for line in instruction_text:
                instruction_canvas.create_text(10, y_offset, anchor="nw", text=line, fill="black", font=("Arial", 10))
                y_offset += 20

        def hide_instructions(event):
            instruction_canvas.delete("all")
            instruction_canvas.place_forget()

        instruction_button.bind("<Enter>", show_instructions)
        instruction_button.bind("<Leave>", hide_instructions)

        entry_xstru_excelfile = Entry(canvas,bd=0, bg="#FFFFFF", fg="#000716", highlightthickness=0)
        canvas.create_window(114.0, 326.0, anchor="nw", window=entry_xstru_excelfile, width=300.0, height=18.0)

        load_xstru_excelfile_image = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_5.png"))
        canvas.button_image_5 = load_xstru_excelfile_image
        load_xstru_excelfile = tk.Button(
            canvas,
            image=load_xstru_excelfile_image,
            borderwidth=0,
            highlightthickness=0,
            command=on_load_beams,
            bg="green",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(43, 327, anchor="nw", window=load_xstru_excelfile, width=54.0, height=20.0)

        def browse_excel_file():
            file_path = filedialog.askopenfilename(
                title="Select XSTRUCT File",
                filetypes=[("XSTRUCT Files", "*.xlsx")]
            )
            if file_path:
                entry_xstru_excelfile.delete(0, "end")
                entry_xstru_excelfile.insert(0, file_path)

        entry_xstru_excelfile.bind("<Button-1>", lambda e: browse_excel_file())

        def browse_staad_file():
            detect_open_staad_file()
            if not entry_load_staad.get() or entry_load_staad.get() == "No STAAD file detected":
                file_path = filedialog.askopenfilename(
                    title="Select STAAD File",
                    filetypes=[("STAAD Files", "*.std")]
                )
                if file_path:
                    entry_load_staad.delete(0, "end")
                    entry_load_staad.insert(0, file_path)

        def check_staad_open():
            try:
                root = Root()
                current_file = root.GetSTAADFile()
                if current_file:
                    return current_file
                else:
                    return
            except Exception as e:
                pass
                return

        def detect_open_staad_file():
            staad_file = check_staad_open()
            if staad_file:
                entry_load_staad.delete(0, "end")
                entry_load_staad.insert(0, staad_file)
            else:
                entry_load_staad.delete(0, "end")
                entry_load_staad.insert(0, "No STAAD file detected")

        class BeamDataFetcher:
            def __init__(self):
                try:
                    self.geometry = Geometry()
                    self.properties = Properties()
                    self.output = Output()
                except Exception as e:
                    logging.error("Error initializing OpenSTAAD: STAAD is not currently open.")
                    raise RuntimeError("Error initializing OpenSTAAD: STAAD is not currently open.") from e

            def fetch_all_beam_data(self, start_load_case, end_load_case):
                beam_list = self.geometry.GetBeamList()
                if not beam_list:
                    logging.warning("No beams found in the model.")
                    return {}
                all_beams_data = {}
                for beam_number in beam_list:
                    beam_data = self.get_beam_data(beam_number, start_load_case, end_load_case)
                    if beam_data:
                        all_beams_data[beam_number] = beam_data
                    else:
                        logging.warning(f"Skipping Beam {beam_number} due to invalid data.")

                return all_beams_data

            def fetch_all_node_data(self):
                node_list = self.geometry.GetNodeList()
                node_data = {}
                for node_number in node_list:
                    coordinates = self.geometry.GetNodeCoordinates(node_number)
                    if coordinates:
                        node_data[node_number] = coordinates
                    else:
                        logging.warning(f"Warning: No coordinates found for Node {node_number}")
                return node_data

            def get_beam_data(self, beam_number, start_load_case, end_load_case):
                try:
                    length = self.geometry.GetBeamLength(beam_number)
                    if length is None:
                        raise ValueError(f"Invalid beam length for Beam {beam_number}")

                    section_name = self.properties.GetBeamSectionName(beam_number) or "Unknown Section"
                    section_ref_no = self.properties.GetBeamSectionPropertyRefNo(beam_number)
                    section_properties = self.properties.GetSectionPropertyValues(section_ref_no)
                    alpha_angle = self.properties.GetAlphaAngleForSection(section_ref_no)
                    member_release = self.properties.GetMemberReleaseSpecEx(beam_number)
                    member_spec_code = self.properties.GetMemberSpecCode(beam_number) or "0"
                    start_node_id, end_node_id = self.geometry.GetMemberIncidence(beam_number)
                    start_coordinates = self.geometry.GetNodeCoordinates(start_node_id)
                    end_coordinates = self.geometry.GetNodeCoordinates(end_node_id)

                    if not (start_coordinates and end_coordinates):
                        raise ValueError(f"Invalid node coordinates for Beam {beam_number}")
                    forces_start, forces_end = self.initialize_forces()
                    for load_case in range(start_load_case, end_load_case + 1):
                        start_forces, end_forces = self.fetch_beam_end_forces(beam_number, load_case)
                        self.update_forces(forces_start, start_forces)
                        self.update_forces(forces_end, end_forces)

                    return {
                        'length': round(length, 4),
                        'section_name': section_name,
                        'section_properties': section_properties,
                        'alpha_angle': alpha_angle,
                        'member_release': member_release,
                        'member_spec_code': member_spec_code,
                        'start_coordinates': start_coordinates,
                        'end_coordinates': end_coordinates,
                        'start_forces': forces_start,
                        'end_forces': forces_end
                    }
                except Exception as e:
                    logging.error(f"Error fetching data for Beam {beam_number}: {e}")
                    return None

            def initialize_forces(self):
                return ({force_type: {'positive': None, 'negative': None} for force_type in
                         ['Axial Force', 'Shear-Y', 'Shear-Z', 'Torsion', 'Moment-Y', 'Moment-Z']},
                        {force_type: {'positive': None, 'negative': None} for force_type in
                         ['Axial Force', 'Shear-Y', 'Shear-Z', 'Torsion', 'Moment-Y', 'Moment-Z']})

            def fetch_beam_end_forces(self, beam_number, load_case):
                try:
                    forces_start = self.output.GetMemberEndForces(beam=beam_number, start=True, lc=load_case)
                    forces_end = self.output.GetMemberEndForces(beam=beam_number, start=False, lc=load_case)
                    return self.parse_forces(forces_start), self.parse_forces(forces_end)
                except Exception as e:
                    logging.error(f"Error fetching forces for Beam {beam_number}, Load Case {load_case}: {e}")
                    return None, None

            def parse_forces(self, forces):
                return {force_type: round(forces[i], 4) for i, force_type in enumerate(
                    ['Axial Force', 'Shear-Y', 'Shear-Z', 'Torsion', 'Moment-Y', 'Moment-Z'])}

            def update_forces(self, forces_dict, new_forces):
                if new_forces:
                    for force_type, values in new_forces.items():
                        if values:
                            if forces_dict[force_type]['positive'] is None or values > forces_dict[force_type][
                                'positive']:
                                forces_dict[force_type]['positive'] = values
                            if forces_dict[force_type]['negative'] is None or values < forces_dict[force_type][
                                'negative']:
                                forces_dict[force_type]['negative'] = values

        def save_beam_data_to_excel(beams_data, file_path, node_data):
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Beam Data"
            headers = ["Beam Number", "Length", "Section Name", "Alpha Angle", "Member Release",
                       "Member Spec Code", "Start X", "Start Y", "Start Z", "End X", "End Y", "End Z",
                       "Start Axial Force", "Start Shear-Y", "Start Shear-Z", "Start Torsion", "Start Moment-Y",
                       "Start Moment-Z", "End Axial Force", "End Shear-Y", "End Shear-Z", "End Torsion",
                       "End Moment-Y", "End Moment-Z"]
            sheet.append(headers)
            for beam_number, beam_data in beams_data.items():
                row = [str(beam_number), str(beam_data.get('length', '0')), beam_data.get('section_name', 'N/A'),
                       str(beam_data.get('alpha_angle', '0')), str(beam_data.get('member_release', 'None')),
                       beam_data.get('member_spec_code', '0')] + list(
                    beam_data['start_coordinates'] or (None, None, None)) + \
                      list(beam_data['end_coordinates'] or (None, None, None))
                start_forces = {k: beam_data['start_forces'][k]['positive'] or 0 for k in beam_data['start_forces']}
                end_forces = {k: beam_data['end_forces'][k]['positive'] or 0 for k in beam_data['end_forces']}
                row.extend([start_forces['Axial Force'], start_forces['Shear-Y'], start_forces['Shear-Z'],
                            start_forces['Torsion'], start_forces['Moment-Y'], start_forces['Moment-Z']])
                row.extend([end_forces['Axial Force'], end_forces['Shear-Y'], end_forces['Shear-Z'],
                            end_forces['Torsion'], end_forces['Moment-Y'], end_forces['Moment-Z']])
                sheet.append(row)
            try:
                workbook.save(file_path)
                messagebox.showinfo("File Saved", f"Excel file saved at: {file_path}")
            except Exception as e:
                logging.error(f"Error saving Excel file: {e}")

        def fetch_and_save_data_with_popup():
            popup = tk.Toplevel(self)
            popup.title("Load Case Input")
            width, height = 300, 200
            popup.geometry(f"{width}x{height}")
            popup.update_idletasks()
            screen_width = popup.winfo_screenwidth()
            screen_height = popup.winfo_screenheight()
            x = (screen_width // 2) - (width // 2)
            y = (screen_height // 2) - (height // 2)
            popup.geometry(f"+{x}+{y}")

            image_path = self.get_resource_path("images_staadpage/image_1.png")
            image_image_1 = PhotoImage(file=image_path)
            popup.background_label = Label(popup, image=image_image_1)
            popup.background_label.place(x=0, y=0, relwidth=1, relheight=1)
            popup.background_label.image = image_image_1
            Label(popup, text="Enter Start Load Case:", bg="#2596be").pack(
                pady=(20, 5))  # Adjust padding for visibility
            start_load_case_var = IntVar()
            start_load_case_entry = Entry(popup, textvariable=start_load_case_var)
            start_load_case_entry.pack()
            Label(popup, text="Enter End Load Case:", bg="#2596be").pack(pady=5)
            end_load_case_var = IntVar()
            end_load_case_entry = Entry(popup, textvariable=end_load_case_var)
            end_load_case_entry.pack()

            def on_confirm():
                try:
                    start_load_case = start_load_case_var.get()
                    end_load_case = end_load_case_var.get()
                    file_path = entry_xstru_excelfile.get()
                    if not file_path:
                        raise ValueError("Please specify a file path for the Excel file.")
                    fetcher = BeamDataFetcher()
                    node_data = fetcher.fetch_all_node_data()  # Get all node data
                    beam_data = fetcher.fetch_all_beam_data(start_load_case, end_load_case)  # Get all beam data

                    if not beam_data:
                        messagebox.showerror("Error", "No beam data available to save.")
                        return
                    save_beam_data_to_excel(beam_data, file_path, node_data)  # Pass node_data
                    messagebox.showinfo("Success", f"Data successfully saved to {file_path}")
                except Exception as e:
                    messagebox.showerror("Error", f"Error fetching or saving data: {e}")
                finally:
                    popup.destroy()

            confirm_button = Button(popup, text="Confirm", command=on_confirm, bg="green", fg="white")
            confirm_button.pack(pady=10)

        entry_load_staad = Entry(canvas, bd=0, bg="#FFFFFF", fg="#000716", highlightthickness=0)
        canvas.create_window(114.0, 454.0, anchor="nw", window=entry_load_staad, width=300.0, height=18.0)
        entry_load_staad.bind("<Button-1>", lambda e: browse_staad_file())
        detect_open_staad_file()

        load_staad_data = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_6.png"))
        canvas.button_image_6 = load_staad_data
        load_staad_data = tk.Button(
            canvas,
            image=load_staad_data,
            borderwidth=0,
            highlightthickness=0,
            command=fetch_and_save_data_with_popup,
            bg="green",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(43, 454, anchor="nw", window=load_staad_data, width=54.0, height=20.0)

        def append_selected_beams_to_excel():
            geometry = Geometry()
            selected_beams = geometry.GetSelectedBeams()
            if not selected_beams:
                messagebox.showinfo("No Selection", "No beams selected in STAAD.")
                return
            excel_path = entry_xstru_excelfile.get()  # Path from your entry widget
            if not excel_path:
                messagebox.showerror("Error", "Please specify a valid file path for the Excel file.")
                return

            def load_case_popup():
                popup = tk.Toplevel(self)
                popup.title("Load Case Input")
                width, height = 300, 200
                popup.geometry(f"{width}x{height}")

                popup.update_idletasks()
                screen_width = popup.winfo_screenwidth()
                screen_height = popup.winfo_screenheight()
                x = (screen_width // 2) - (width // 2)
                y = (screen_height // 2) - (height // 2)
                popup.geometry(f"+{x}+{y}")

                image_path = self.get_resource_path("images_staadpage/image_1.png")
                image_image_1 = PhotoImage(file=image_path)
                popup.background_label = Label(popup, image=image_image_1)
                popup.background_label.place(x=0, y=0, relwidth=1, relheight=1)
                popup.background_label.image = image_image_1

                Label(popup, text="Enter Start Load Case:", bg="#2596be").pack(
                    pady=(20, 5))
                start_load_case_var = IntVar()
                start_load_case_entry = Entry(popup, textvariable=start_load_case_var)
                start_load_case_entry.pack()
                Label(popup, text="Enter End Load Case:", bg="#2596be").pack(pady=5)
                end_load_case_var = IntVar()
                end_load_case_entry = Entry(popup, textvariable=end_load_case_var)
                end_load_case_entry.pack()

                def on_confirm():
                    try:
                        start_load_case = start_load_case_var.get()
                        end_load_case = end_load_case_var.get()
                        file_path = entry_xstru_excelfile.get()
                        if not file_path:
                            raise ValueError("Please specify a file path for the Excel file.")
                        fetcher = BeamDataFetcher()
                        node_data = fetcher.fetch_all_node_data()
                        beam_data = fetcher.fetch_all_beam_data(start_load_case, end_load_case)
                        if not beam_data:
                            messagebox.showerror("Error", "No beam data available to save.")
                            return
                        save_beam_data_to_excel(beam_data, file_path, node_data)
                        workbook = load_workbook(file_path)
                        sheet = workbook.active
                        selected_beam_ids = set(str(beam_id) for beam_id in selected_beams)
                        rows_to_delete = []
                        for row in sheet.iter_rows(min_row=2):
                            beam_id_cell = row[0]
                            if beam_id_cell and beam_id_cell.value is not None:
                                beam_id = str(beam_id_cell.value)

                                if beam_id not in selected_beam_ids:
                                    rows_to_delete.append(beam_id_cell.row)

                                else:
                                    pass
                        for row_num in sorted(rows_to_delete, reverse=True):
                            sheet.delete_rows(row_num)
                        workbook.save(file_path)
                        workbook.close()
                    except Exception as e:
                        messagebox.showerror("Error", f"Error fetching, saving, or filtering data: {e}")
                    finally:
                        popup.destroy()

                confirm_button = Button(popup, text="Confirm", command=on_confirm, bg="green", fg="white")
                confirm_button.pack(pady=10)

            load_case_popup()

        select_from_staad = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_7.png"))
        canvas.select_from_staad = select_from_staad
        select_from_staad_button = tk.Button(canvas,
            image=select_from_staad,
            borderwidth=0,
            highlightthickness=0,
            command=append_selected_beams_to_excel,
            bg="green",
            activebackground="green",
            relief="flat"
        )

        canvas.create_window(167, 494, anchor="nw", window=select_from_staad_button, width=134.0, height=20.0)

        design_codes = ["NSCP 2015", "NSCP 2025"]
        selected_design_code = StringVar()
        selected_design_code.set(design_codes[0])
        dropdown_design_code = OptionMenu(canvas, selected_design_code, *design_codes)
        dropdown_design_code.config(
            bg="#FFFFFF",  #
            fg="#000716",  #
            font=("Helvetica", 10),
            relief="flat",
            highlightthickness=0,
            width=200
        )
        dropdown_design_code["menu"].config(
            bg="#FFFFFF",
            fg="#000716",
            font=("Helvetica", 10),
            relief="flat",
        )
        canvas.create_window(114.0, 576.0, anchor="nw", window=dropdown_design_code, width=300.0, height=18.0)

        button_image_1 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_1.png"))
        button_1 = tk.Button(canvas,
            image=button_image_1,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: self.switch_indication(self.staad_btn_indicator, self.staad_page),
            relief="flat"
        )
        canvas.button_image_1 = button_image_1
        canvas.create_window(0, 48.0, anchor="nw", window=button_1, width=40.0, height=40.0)

        button_image_1 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_1.png"))
        button_1 = tk.Button(canvas1,
            image=button_image_1,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: self.switch_indication(self.staad_btn_indicator, self.staad_page),
            relief="flat"
        )
        canvas1.button_image_1 = button_image_1
        canvas1.create_window(0, 48.0, anchor="nw", window=button_1, width=40.0, height=40.0)

        def toggle_canvas():
            canvas.place_forget()
            canvas1.place(x=0, y=0)

        button_image_2 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_2.png"))
        canvas.button_image_2 = button_image_2
        button_2 = tk.Button(canvas,
            image=button_image_2,
            borderwidth=0,
            highlightthickness=0,
            command=toggle_canvas,
            relief="flat"
        )
        canvas.create_window(44.0, 48.0, anchor="nw", window=button_2, width=40.0, height=40.0)

        button_image_2 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_2.png"))
        canvas1.button_image_2 = button_image_2
        button_2 = tk.Button(canvas1,
            image=button_image_2,
            borderwidth=0,
            highlightthickness=0,
            command=toggle_canvas,
            relief="flat"
        )
        canvas1.create_window(44.0, 48.0, anchor="nw", window=button_2, width=40.0, height=40.0)

        fb_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_3.png"))
        canvas.button_image_3 = fb_button
        fb_button = tk.Button(canvas,
            image=fb_button,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: webbrowser.open("https://www.facebook.com/@xstructures"),
            relief="flat"
        )
        canvas.create_window(973, 50.0, anchor="nw", window=fb_button, width=35.0, height=35.0)

        fb_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_3.png"))
        canvas1.button_image_3 = fb_button
        fb_button = tk.Button(canvas1,
            image=fb_button,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: webbrowser.open("https://www.facebook.com/@xstructures"),
            relief="flat"
        )
        canvas1.create_window(973, 50.0, anchor="nw", window=fb_button, width=35.0, height=35.0)

        button_image_8 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_8.png"))
        canvas.button_image_8 = button_image_8
        button_8 = tk.Button(canvas,
            image=button_image_8,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: self.switch_indication(self.concrete_btn_indicator, self.concrete_page),
            bg="green",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(169, 616, anchor="nw", window=button_8, width=134.0, height=20.0)

        button_image_9 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_9.png"))
        canvas.button_image_9 = button_image_9
        button_9 = tk.Button(canvas,
            image=button_image_9,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: self.switch_indication(self.steel_btn_indicator, self.steel_page),
            bg="green",
            activebackground="green",
            relief="flat"
        )
        canvas.create_window(183, 648, anchor="nw", window=button_9, width=102.0, height=20.0)

    def concrete_page(self):
        concrete_page_fm = tk.Frame(self.page_frame)
        concrete_page_fm.place(relwidth=1.0, relheight=1.0)

        canvas = Canvas(concrete_page_fm, bg="#FFFFFF", height=720, width=1016, bd=0, highlightthickness=0)
        canvas.place(x=0, y=0)
        canvas.create_rectangle(0, 0, 1016, 720, fill="#FFFFFF", outline="")
        label1 = tk.Label(canvas, text="PART 1 CONCRETE DESIGN", bg="#FFFFFF", font=("Arial", 14))
        canvas.create_window(508, 360, window=label1)

        # Canvas for Part 2
        canvas1 = Canvas(concrete_page_fm, bg="#F5F5F5", height=720, width=1016, bd=0, highlightthickness=0)
        canvas1.create_rectangle(0, 0, 1016, 720, fill="#F5F5F5", outline="")
        label2 = tk.Label(canvas1, text="PART 2 CONCRETE DESIGN", bg="#F5F5F5", font=("Arial", 14))
        canvas1.create_window(508, 360, window=label2)

        image_image_1 = PhotoImage(file=self.get_resource_path("images_staadpage/image_1.png"))
        canvas.image_1 = image_image_1
        canvas.create_image(501.0, 396.0, image=image_image_1)
        canvas1.create_image(501.0, 396.0, image=image_image_1)

        image_image_7 = PhotoImage(file=self.get_resource_path("images_staadpage/image_7.png"))
        canvas.image_image_7 = image_image_7
        canvas.create_image(500.0, 708.0, image=image_image_7)
        canvas1.create_image(500.0, 708.0, image=image_image_7)

        image_image_10 = PhotoImage(file=self.get_resource_path("images_staadpage/image_10.png"))
        canvas.image_image_10 = image_image_10
        canvas.create_image(500, 24.0, image=image_image_10)
        canvas1.create_image(500, 24.0, image=image_image_10)

        button_image_10 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_10.png"))
        canvas.button_image_10 = button_image_10
        button_10 = tk.Button(canvas,
            image=button_image_10,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: self.switch_indication(self.concrete_btn_indicator, self.concrete_page),
            relief="flat"
        )
        canvas.create_window(0, 48.0, anchor="nw", window=button_10, width=40.0, height=40.0)

        button_image_10 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_10.png"))
        canvas1.button_image_10 = button_image_10
        button_10 = tk.Button(canvas1,
            image=button_image_10,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: self.switch_indication(self.concrete_btn_indicator, self.concrete_page),
            relief="flat"
        )
        canvas1.create_window(0, 48.0, anchor="nw", window=button_10, width=40.0, height=40.0)

        def toggle_canvas():
            canvas.place_forget()
            canvas1.place(x=0, y=0)

        button_image_11 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_11.png"))
        canvas.button_image_11 = button_image_11
        button_11 = tk.Button(canvas,
            image=button_image_11,
            borderwidth=0,
            highlightthickness=0,
            command=toggle_canvas,
            relief="flat"
        )
        canvas.create_window(44.0, 48.0, anchor="nw", window=button_11, width=40.0, height=40.0)

        button_image_11 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_11.png"))
        canvas1.button_image_11 = button_image_11
        button_11 = tk.Button(canvas1,
            image=button_image_11,
            borderwidth=0,
            highlightthickness=0,
            command=toggle_canvas,
            relief="flat"
        )
        canvas1.create_window(44.0, 48.0, anchor="nw", window=button_11, width=40.0, height=40.0)

        fb_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_3.png"))
        canvas.button_image_3 = fb_button
        fb_button = tk.Button(canvas,
            image=fb_button,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: webbrowser.open("https://www.facebook.com/@xstructures"),
            relief="flat"
        )
        canvas.create_window(973, 50.0, anchor="nw", window=fb_button, width=35.0, height=35.0)

        fb_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_3.png"))
        canvas1.button_image_3 = fb_button
        fb_button = tk.Button(canvas1,
            image=fb_button,
            borderwidth=0,
            highlightthickness=0,
            command=lambda: webbrowser.open("https://www.facebook.com/@xstructures"),
            relief="flat"
        )
        canvas1.create_window(973, 50.0, anchor="nw", window=fb_button, width=35.0, height=35.0)

    def steel_page(self):
        steel_page_fm = tk.Frame(self.page_frame)
        steel_page_fm.place(relwidth=1.0, relheight=1.0)

        # Canvas 1 (Part 1 Steel Design)
        canvas = Canvas(steel_page_fm, bg="#FFFFFF", height=720, width=1016, bd=0, highlightthickness=0)
        canvas.place(x=0, y=0)
        canvas.create_rectangle(0, 0, 1016, 720, fill="#FFFFFF", outline="")
        label1 = tk.Label(canvas, text="PART 1 STEEL DESIGN", bg="#FFFFFF", font=("Arial", 14))
        canvas.create_window(508, 360, window=label1)

        # Canvas 2 (Part 2 Steel Design)
        canvas1 = Canvas(steel_page_fm, bg="#F5F5F5", height=720, width=1016, bd=0, highlightthickness=0)
        canvas1.create_rectangle(0, 0, 1016, 720, fill="#F5F5F5", outline="")
        label2 = tk.Label(canvas1, text="PART 2 STEEL DESIGN", bg="#F5F5F5", font=("Arial", 14))
        canvas1.create_window(508, 360, window=label2)

        image_image_1 = PhotoImage(file=self.get_resource_path("images_staadpage/image_1.png"))
        canvas.image_1 = image_image_1
        canvas.create_image(501.0, 396.0, image=image_image_1)
        canvas1.create_image(501.0, 396.0, image=image_image_1)

        image_image_7 = PhotoImage(file=self.get_resource_path("images_staadpage/image_7.png"))
        canvas.image_image_7 = image_image_7
        canvas.create_image(500.0, 708.0, image=image_image_7)
        canvas1.create_image(500.0, 708.0, image=image_image_7)

        image_image_11 = PhotoImage(file=self.get_resource_path("images_staadpage/image_11.png"))
        canvas.image_image_10 = image_image_11
        canvas.create_image(500, 24.0, image=image_image_11)
        canvas1.create_image(500, 24.0, image=image_image_11)

        button_image_10 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_10.png"))
        canvas.button_image_10 = button_image_10
        button_10 = tk.Button(canvas,
                              image=button_image_10,
                              borderwidth=0,
                              highlightthickness=0,
                              command=lambda: self.switch_indication(self.steel_btn_indicator, self.steel_page),
                              relief="flat"
                              )
        canvas.create_window(0, 48.0, anchor="nw", window=button_10, width=40.0, height=40.0)

        button_image_10 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_10.png"))
        canvas1.button_image_10 = button_image_10
        button_10 = tk.Button(canvas1,
                              image=button_image_10,
                              borderwidth=0,
                              highlightthickness=0,
                              command=lambda: self.switch_indication(self.steel_btn_indicator, self.steel_page),
                              relief="flat"
                              )
        canvas1.create_window(0, 48.0, anchor="nw", window=button_10, width=40.0, height=40.0)

        def toggle_canvas():
            canvas.place_forget()
            canvas1.place(x=0, y=0)

        button_image_11 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_11.png"))
        canvas.button_image_11 = button_image_11
        button_11 = tk.Button(canvas,
                              image=button_image_11,
                              borderwidth=0,
                              highlightthickness=0,
                              command=toggle_canvas,
                              relief="flat"
                              )
        canvas.create_window(44.0, 48.0, anchor="nw", window=button_11, width=40.0, height=40.0)

        button_image_11 = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_11.png"))
        canvas1.button_image_11 = button_image_11
        button_11 = tk.Button(canvas1,
                              image=button_image_11,
                              borderwidth=0,
                              highlightthickness=0,
                              command=toggle_canvas,
                              relief="flat"
                              )
        canvas1.create_window(44.0, 48.0, anchor="nw", window=button_11, width=40.0, height=40.0)

        fb_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_3.png"))
        canvas.button_image_3 = fb_button
        fb_button = tk.Button(canvas,
                              image=fb_button,
                              borderwidth=0,
                              highlightthickness=0,
                              command=lambda: webbrowser.open("https://www.facebook.com/@xstructures"),
                              relief="flat"
                              )
        canvas.create_window(973, 50.0, anchor="nw", window=fb_button, width=35.0, height=35.0)

        fb_button = tk.PhotoImage(file=self.get_resource_path("images_staadpage/button_3.png"))
        canvas1.button_image_3 = fb_button
        fb_button = tk.Button(canvas1,
                              image=fb_button,
                              borderwidth=0,
                              highlightthickness=0,
                              command=lambda: webbrowser.open("https://www.facebook.com/@xstructures"),
                              relief="flat"
                              )
        canvas1.create_window(973, 50.0, anchor="nw", window=fb_button, width=35.0, height=35.0)

    def extend_menu_bar(self):
        """Extend the menu bar."""
        self.extending_animation()
        self.menu_bar_frame.lift()
        self.toggle_menu_btn.config(image=self.close_icon, command=self.fold_menu_bar)

    def fold_menu_bar(self):
        """Fold the menu bar."""
        self.folding_animation()
        self.toggle_menu_btn.config(image=self.toggle_icon, command=self.extend_menu_bar)

    def extending_animation(self):
        """Animation for extending the menu bar."""
        current_width = self.menu_bar_frame.winfo_width()
        if current_width < 200:
            current_width += 10
            self.menu_bar_frame.config(width=current_width)
            self.after(5, self.extending_animation)

    def folding_animation(self):
        """Animation for folding the menu bar."""
        current_width = self.menu_bar_frame.winfo_width()
        if current_width > 60:
            current_width -= 10
            self.menu_bar_frame.config(width=current_width)
            self.after(5, self.folding_animation)

    def show_loading_screen(self):
        """Show a loading screen."""
        self.withdraw()

        splash = tk.Toplevel(self)
        splash.geometry("288x480")
        splash.overrideredirect(True)

        # Center the splash screen
        screen_width = self.winfo_screenwidth()
        screen_height = self.winfo_screenheight()
        position_top = int(screen_height / 2 - 480 / 2)
        position_left = int(screen_width / 2 - 288 / 2)
        splash.geometry(f"288x480+{position_left}+{position_top}")

        # Loading animation
        gif_path = self.get_resource_path("images_sidebar/loading.gif")

        gif_image = Image.open(gif_path)
        gif_frames = [ImageTk.PhotoImage(frame) for frame in ImageSequence.Iterator(gif_image)]
        gif_label = tk.Label(splash)
        gif_label.pack()

        def animate(frame_index=0):
            if gif_label.winfo_exists():
                gif_label.config(image=gif_frames[frame_index])
                frame_index = (frame_index + 1) % len(gif_frames)
                splash.after(100, animate, frame_index)

        animate()

        # Start a thread to simulate a loading task
        threading.Thread(target=self.run_main_task, args=(splash,)).start()

    def run_main_task(self, splash):
        """Simulate a main task."""
        time.sleep(2)  # Simulate loading
        splash.destroy()
        self.destroy()  # Close this window
        self.master.deiconify()

    def get_resource_path(self, relative_path):
        """
        Returns the absolute path to a resource, handling both frozen and non-frozen environments.
        """
        if getattr(sys, 'frozen', False):
            base_path = sys._MEIPASS  # If running as a bundled app (PyInstaller)
        else:
            base_path = os.path.abspath(".")  # If running as a normal Python script
        return os.path.join(base_path, relative_path)


# Main Application
if __name__ == "__main__":
    root = tk.Tk()
    root.withdraw()  # Hide the main Tk window if using Toplevel as the primary interface
    app = Interface(root)
    root.mainloop()
